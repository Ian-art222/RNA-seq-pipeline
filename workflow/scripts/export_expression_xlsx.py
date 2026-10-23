#!/usr/bin/env python3
"""Validate related CSV tables and preserve literal identifiers in Excel."""
import argparse
import csv
import math
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from build_expression_matrices import checked_ids, finite_nonnegative, nonnegative_integer, require_same_ids, validate_tpm


def read_table(path, identifier, integer=False):
    with Path(path).open(encoding='utf-8', newline='') as handle:
        reader = csv.reader(handle)
        header = next(reader, None)
        if not header or header[0] != identifier:
            raise ValueError(f'{path}: expected first column {identifier}')
        checked_ids(header, f'{path}: columns')
        if len(header) < 2:
            raise ValueError(f'{path}: no data columns')
        rows, seen = [], set()
        for number, row in enumerate(reader, 2):
            if len(row) != len(header) or not row[0] or row[0] in seen:
                raise ValueError(f'{path}:{number}: invalid width, missing or duplicate {identifier}')
            seen.add(row[0])
            values = [row[0]]
            for column, value in zip(header[1:], row[1:]):
                label = f'{path}:{number}/{column}'
                if integer:
                    values.append(nonnegative_integer(value, label))
                else:
                    numeric = finite_nonnegative(value, label)
                    values.append(int(numeric) if numeric.is_integer() else numeric)
            rows.append(values)
        if not rows:
            raise ValueError(f'{path}: no data rows')
    return header, rows


def validate_tables(alignment, count, fpkm, tpm):
    sample_names = count[0][1:]
    genes = [row[0] for row in count[1]]
    for label, table in (('FPKM', fpkm), ('TPM', tpm)):
        if table[0] != count[0]:
            raise ValueError(f'{label}: sample IDs/order differ from the count matrix')
        if [row[0] for row in table[1]] != genes:
            raise ValueError(f'{label}: gene IDs/order differ from the count matrix')
    require_same_ids([row[0] for row in alignment[1]], sample_names, 'Alignment report versus expression samples')
    for row in count[1]:
        if any(not isinstance(value, int) for value in row[1:]):
            raise ValueError(f'Counts must be nonnegative integers: {row[0]}')
    for index, sample in enumerate(sample_names, 1):
        validate_tpm((row[index] for row in tpm[1]), (row[index] for row in count[1]),
                     (row[index] for row in fpkm[1]), sample)


def write_workbook(output_path, tables):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.' + output_path.stem + '.', suffix='.xlsx', dir=output_path.parent)
    os.close(fd)
    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, (header, rows) in tables.items():
            if len(header) > 16384 or len(rows) + 1 > 1048576:
                raise ValueError(f'{name}: table exceeds Excel worksheet row/column limits; use the CSV matrices')
            sheet = workbook.create_sheet(name)
            sheet.append(header)
            for cell in sheet[1]:
                cell.data_type = 's'
            for index, row in enumerate(rows, 2):
                sheet.append(row)
                # Explicit string cells also preserve IDs beginning '=' literally.
                sheet.cell(index, 1).data_type = 's'
            sheet.freeze_panes = 'B2'
            sheet.auto_filter.ref = sheet.dimensions
        workbook.save(temporary)
        workbook.close()
        reopened = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if reopened.sheetnames != list(tables):
                raise ValueError('Workbook sheet validation failed')
            for name, (header, rows) in tables.items():
                sheet = reopened[name]
                if (sheet.max_row, sheet.max_column) != (len(rows) + 1, len(header)):
                    raise ValueError(f'{name}: workbook dimensions changed')
                actual_rows = sheet.iter_rows(values_only=True)
                if list(next(actual_rows)) != header:
                    raise ValueError(f'{name}: workbook column identifiers changed')
                for actual, expected in zip(actual_rows, rows):
                    if actual[0] != expected[0] or not isinstance(actual[0], str):
                        raise ValueError(f'{name}: workbook row identifier changed')
                    for actual_value, expected_value in zip(actual[1:], expected[1:]):
                        numeric_match = (actual_value == expected_value if isinstance(expected_value, int)
                                         else isinstance(actual_value, (int, float)) and
                                         math.isclose(actual_value, expected_value, rel_tol=1e-12, abs_tol=1e-12))
                        if not isinstance(actual_value, (int, float)) or not numeric_match:
                            raise ValueError(f'{name}: workbook numeric value changed')
        finally:
            reopened.close()
        os.replace(temporary, output_path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--alignment-report', required=True)
    parser.add_argument('--count', required=True)
    parser.add_argument('--fpkm', required=True)
    parser.add_argument('--tpm', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    alignment = read_table(args.alignment_report, 'SampleID')
    count = read_table(args.count, 'GeneID', integer=True)
    fpkm = read_table(args.fpkm, 'GeneID')
    tpm = read_table(args.tpm, 'GeneID')
    validate_tables(alignment, count, fpkm, tpm)
    write_workbook(args.output, {'alignment_report': alignment, 'gene_count': count, 'gene_fpkm': fpkm, 'gene_tpm': tpm})
    print(f'[INFO] Validated and wrote Excel workbook: {args.output}')


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError) as error:
        raise SystemExit(f'[ERROR] {error}')
